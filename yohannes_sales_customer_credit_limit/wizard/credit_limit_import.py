import base64
import io
import logging

try:
    import openpyxl
except ImportError:
    openpyxl = None

from odoo import models, fields, api, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

class CreditLimitImport(models.TransientModel):
    _name = 'yohannes.credit.limit.import.wizard'
    _description = 'Import Customer Credit Limits'

    file = fields.Binary(string='Excel File', required=False)
    file_name = fields.Char(string='File Name')
    import_log = fields.Text(string="Import Log", readonly=True)
    
    def action_download_template(self):
        if not openpyxl:
            raise UserError(_("Python library 'openpyxl' is required to generate the template. Please install it on the server."))
            
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Credit Limits"
        
        # Headers to match UI exactly
        headers = ["Name", "Tin No", "City/sub-city", "Payment terms", "Credit limit"]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            
        # Sample Data
        ws.cell(row=2, column=1, value="Sample Customer")
        ws.cell(row=2, column=2, value="0001234567")
        ws.cell(row=2, column=3, value="Addis Ababa")
        ws.cell(row=2, column=4, value="30 Days")
        ws.cell(row=2, column=5, value=15000)
        
        # Save to in-memory byte stream
        fp = io.BytesIO()
        wb.save(fp)
        fp.seek(0)
        file_data = base64.b64encode(fp.read())

        attachment = self.env['ir.attachment'].create({
            'name': 'Credit_Limit_Import_Template.xlsx',
            'type': 'binary',
            'datas': file_data,
            'res_model': 'yohannes.credit.limit.import.wizard',
            'public': True,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def action_import(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_("Please upload an Excel file."))
            
        if not openpyxl:
            raise UserError(_("Python library 'openpyxl' is required to parse Excel files. Please install it on the server."))

        try:
            file_data = base64.b64decode(self.file)
            fp = io.BytesIO(file_data)
            wb = openpyxl.load_workbook(fp, data_only=True)
            ws = wb.active
        except Exception as e:
            raise UserError(_("Invalid file format. Please upload a valid .xlsx file. Details: %s") % str(e))

        log_messages = []
        customer_credit_env = self.env['customer.credit']
        partner_env = self.env['res.partner']
        payment_term_env = self.env['account.payment.term']
        
        success_count = 0
        update_count = 0
        error_count = 0

        # Assuming first row is header
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # Read columns based on new template: 
            # 0: Name, 1: Tin No, 2: City/sub-city, 3: Payment terms, 4: Credit limit
            if not row or not any(row):
                continue # Skip empty rows

            tin = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
            term_str = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ''
            limit_val = row[4] if len(row) > 4 else None

            if not tin:
                # If TIN is missing, try fallback to Name
                customer_name = str(row[0]).strip() if row[0] is not None else ''
                if customer_name:
                    partners = partner_env.search([('name', '=ilike', customer_name)])
                    if not partners:
                        log_messages.append(f"Row {row_idx}: Skipped. TIN is missing and Customer Name '{customer_name}' was not found.")
                        error_count += 1
                        continue
                else:    
                    log_messages.append(f"Row {row_idx}: Skipped. Both TIN and Name missing.")
                    error_count += 1
                    continue
            else:
                partners = partner_env.search([('vat', '=', tin)])
                if not partners:
                    log_messages.append(f"Row {row_idx}: Failed. No customer found with Tin No {tin}.")
                    error_count += 1
                    continue

            try:
                credit_limit = float(limit_val) if limit_val is not None else 0.0
            except (ValueError, TypeError):
                log_messages.append(f"Row {row_idx}: Skipped. Invalid Credit Limit value '{limit_val}'.")
                error_count += 1
                continue
                
            partner = partners[0] # Take first match if duplicates exist

            # Find Payment Term if provided
            term_id = False
            if term_str:
                terms = payment_term_env.search([('name', '=ilike', term_str)], limit=1)
                if terms:
                    term_id = terms.id
                else:
                    log_messages.append(f"Row {row_idx}: Warning. Payment term '{term_str}' not found. Leaving blank.")

            # Check if customer already has a credit record
            existing_credit = customer_credit_env.search([('partner_id', '=', partner.id)], limit=1)
            
            vals = {
                'maximum_credit': credit_limit,
            }
            if term_id:
                vals['payment_term_credit'] = term_id

            if existing_credit:
                existing_credit.write(vals)
                log_messages.append(f"Row {row_idx}: Updated existing credit limit for {partner.name}.")
                update_count += 1
            else:
                vals['partner_id'] = partner.id
                # Creating brand new
                customer_credit_env.create(vals)
                log_messages.append(f"Row {row_idx}: Created new credit limit for {partner.name}.")
                success_count += 1

        self.import_log = f"Import Result: {success_count} Created, {update_count} Updated, {error_count} Errors.\n\n" + "\n".join(log_messages)
        
        return {
            'type': 'ir.actions.act_window',
            'name': 'Import Results',
            'res_model': 'yohannes.credit.limit.import.wizard',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }
